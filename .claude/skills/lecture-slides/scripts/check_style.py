"""pptx의 슬라이드 글과 발표자 노트에서 명사형 종결이 아닌 문장을 찾아냄.

python check_style.py <deck.pptx>

'~다', '~요', '~까', '~죠', 마침표로 끝나는 줄을 표시. 기계적 검사라 '바다' 같은 명사도 걸릴 수 있으니
목록을 보고 판단해서 고칠 것
"""
import re, sys
from pptx import Presentation

BAD = re.compile(r"(다|요|까|죠|니다|세요)\s*[.!?]?\s*$")
BAD_BODY = re.compile(BAD.pattern + r"|[^.][.]\s*$")   # 슬라이드 본문은 마침표도 생략 (노트는 문장 구분용 마침표 허용)
prs = Presentation(sys.argv[1])
hits = 0
for i, s in enumerate(prs.slides, 1):
    texts = [("본문", sh.text_frame.text) for sh in s.shapes if sh.has_text_frame]
    for sh in s.shapes:
        if sh.has_table:
            texts += [("표", c.text) for r in sh.table.rows for c in r.cells]
    if s.has_notes_slide:
        texts += [("노트", t) for t in re.split(r"(?<=[.!?])\s+|\n", s.notes_slide.notes_text_frame.text)]
    for where, t in texts:
        for line in t.splitlines():
            line = line.strip()
            pat = BAD if where == "노트" else BAD_BODY
            if line and not re.fullmatch(r"[\d.\s=+−×-]+", line) and pat.search(line):
                hits += 1; print(f"slide {i} [{where}] {line}")
print("의심 문장", hits, "개" if hits else "개 — 문체 통과")
