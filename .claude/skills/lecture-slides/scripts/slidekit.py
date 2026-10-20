"""강의 슬라이드 공통 도구: 디자인 상수, 그림 저장, python-pptx 배치 함수, 애니메이션 표시.

사용법 (단원별 deck.py 맨 위):
    import sys; sys.path.insert(0, r"<스킬 폴더>/scripts")
    from slidekit import Deck, mpl_setup, C
    plt = mpl_setup()
    deck = Deck(tag="04 정규분포")

애니메이션은 여기서 도형 이름에만 표시(A순번|효과|시점)하고,
실제 효과는 저장 후 apply_anims.py(PowerPoint)가 붙인다.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- 색 (그래프와 슬라이드가 같은 색을 쓰도록 한곳에서 관리) ----------
class C:
    BLUE, ORANGE, GREEN, GRAY, RED = "#2563EB", "#F97316", "#16A34A", "#64748B", "#DC2626"
    NAVY_HEX, LIGHT_HEX = "#1E293B", "#EFF6FF"

def rgb(hex_):
    hex_ = hex_.lstrip("#")
    return RGBColor(int(hex_[0:2], 16), int(hex_[2:4], 16), int(hex_[4:6], 16))

NAVY, BL, OR, LIGHT, MUTE, WHITE = (rgb(C.NAVY_HEX), rgb(C.BLUE), rgb(C.ORANGE),
                                   rgb(C.LIGHT_HEX), rgb(C.GRAY), RGBColor(255, 255, 255))
FONT = "맑은 고딕"


def mpl_setup(font_size=15):
    """한글 폰트·수식 글꼴·ffmpeg 경로를 맞춘 matplotlib 반환"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rc = {"font.family": "Malgun Gothic", "axes.unicode_minus": False, "font.size": font_size,
          "axes.spines.top": False, "axes.spines.right": False, "mathtext.fontset": "cm"}
    try:
        import imageio_ffmpeg
        rc["animation.ffmpeg_path"] = imageio_ffmpeg.get_ffmpeg_exe()
    except ImportError:
        pass
    plt.rcParams.update(rc)
    return plt


def save_fig(fig, path, tight=True):
    """투명 배경 PNG 저장. 겹쳐서 등장시킬 레이어 그림은 tight=False로 좌표를 고정"""
    import matplotlib.pyplot as plt
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    fig.savefig(path, dpi=200, transparent=True, **({"bbox_inches": "tight"} if tight else {}))
    plt.close(fig)
    return path


def formula_png(tex, path, size=34):
    """수식(LaTeX 문법) → PNG. 슬라이드 글자로 수식을 쓰면 분수·지수가 깨지므로 그림으로 넣음"""
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(6, 1))
    fig.text(0, .5, tex, fontsize=size, va="center")
    return save_fig(fig, path)


def trim_edges(path, px=5):
    """배경을 지운 PNG 사진 가장자리에 남은 가는 선(어두운 배경에서 눈에 띔) 잘라내기. 새 파일 경로 반환"""
    from PIL import Image
    im = Image.open(path); w, h = im.size
    out = os.path.splitext(path)[0] + "_trim.png"
    im.crop((px, px, w - px, h - px)).save(out)
    return out


def load_credits(media_dir):
    """commons.py get 이 기록한 출처 문구 {파일명: 문구}"""
    import json
    p = os.path.join(media_dir, "credits.json")
    return {k: v["credit"] for k, v in json.load(open(p, encoding="utf-8")).items()} if os.path.exists(p) else {}


class Deck:
    def __init__(self, tag="", template=None):
        self.prs = Presentation(template) if template else Presentation()
        if not template:
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6 if len(self.prs.slide_layouts) > 6 else -1]
        self.tag = tag
        self.credits = []

    # ---------- 슬라이드 ----------
    def cover(self, subject, title, subtitle, image=None, credit=None, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
        self.text(s, .9, 1.9, 6.5, 1, subject, 22, rgb("#93C5FD"))
        self.text(s, .9, 2.6, 7, 1.5, title, 60, WHITE, True)
        self.text(s, .9, 4.0, 7, .8, subtitle, 22, rgb("#CBD5E1"))
        if image:
            self.pic(s, image, 8.6, 1.5, h=4.2)
        if credit:
            self.text(s, 8.0, 5.9, 5, .3, credit, 10, rgb("#94A3B8")); self.credits.append(credit)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def slide(self, title, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.18), Inches(1.15))
        bar.fill.solid(); bar.fill.fore_color.rgb = BL; bar.line.fill.background()
        self.text(s, .5, .3, 10, .8, title, 32, NAVY, True)
        self.text(s, 10.3, .38, 2.7, .5, self.tag, 14, MUTE, align=PP_ALIGN.RIGHT)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def credits_slide(self, extra=("그래프·애니메이션: 직접 제작 (Python matplotlib)",)):
        s = self.slide("그림 출처")
        self.text(s, .8, 1.6, 12, 5, sorted(set(self.credits)) + list(extra), 16, MUTE)
        return s

    def save(self, path):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        print("saved", path)

    # ---------- 배치 (단위: 인치, 슬라이드 13.333 × 7.5) ----------
    def text(self, s, l, t, w, h, lines, size=24, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
        tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        for i, ln in enumerate(lines if isinstance(lines, (list, tuple)) else [lines]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(10)
            r = p.add_run(); r.text = ln
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT; r.font.color.rgb = color
        return tb

    def box(self, s, l, t, w, h, fill=LIGHT, line=BL):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        b.adjustments[0] = .06; b.fill.solid(); b.fill.fore_color.rgb = fill
        b.line.color.rgb = line; b.line.width = Pt(2); b.shadow.inherit = False
        return b

    def boxed(self, s, l, t, w, h, runs, size=21, align=PP_ALIGN.CENTER):
        """글자가 들어 있는 상자 (상자째로 애니메이션 걸기 좋음).
        runs: 문자열 또는 [(글자, 색, 굵게), ...] 한 줄"""
        b = self.box(s, l, t, w, h)
        tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(.3)
        p = tf.paragraphs[0]; p.alignment = align
        for txt, color, bold in ([(runs, NAVY, True)] if isinstance(runs, str) else runs):
            r = p.add_run(); r.text = txt
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT; r.font.color.rgb = color
        return b

    def pic(self, s, path, l, t, w=None, h=None):
        return s.shapes.add_picture(path, Inches(l), Inches(t),
                                    Inches(w) if w else None, Inches(h) if h else None)

    def photo(self, s, path, l, t, h=None, credit="", w=None):
        """인터넷 사진 + 바로 아래 작은 출처. 출처는 마지막 '그림 출처' 슬라이드에도 모임"""
        p = self.pic(s, path, l, t, w=w, h=None if w else h)
        if credit:
            pw = Emu(p.width).inches
            self.text(s, l, t + Emu(p.height).inches + .02, max(pw, 3), .3, credit, 10, MUTE)
            self.credits.append(credit)
        return p

    def table(self, s, l, t, w, h, cells, highlight=(), size=18):
        rows, cols = len(cells), len(cells[0])
        gf = s.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(h))
        for r in range(rows):
            for c in range(cols):
                cell = gf.table.cell(r, c); cell.text = cells[r][c]
                p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                hi, head = (r, c) in highlight, r == 0 or c == 0
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb("#FED7AA") if hi else (LIGHT if head else WHITE)
                for run in p.runs:
                    run.font.size, run.font.name, run.font.bold = Pt(size), FONT, hi or head
                    run.font.color.rgb = NAVY
        return gf

    # ---------- 애니메이션 표시 ----------
    def anim(self, s, shape, effect="fade", trigger="click"):
        """effect: fade | wipe(왼쪽부터) | appear,  trigger: click | with | after"""
        s._n = getattr(s, "_n", 0) + 1
        shape.name = f"A{s._n:02d}|{effect}|{trigger}"
        return shape

    def video(self, s, filename, l, t, w, h, trigger="after"):
        """영상 자리 표시. apply_anims.py가 media 폴더의 filename으로 교체하고 재생 효과를 붙임.
        trigger='after'면 슬라이드에 들어가자마자 자동 재생"""
        ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        return self.anim(s, ph, "video:" + filename, trigger)
